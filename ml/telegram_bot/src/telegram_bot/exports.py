"""Выгрузка событий (макет 10): Excel с листами «События» и «По машинам» и тот же список в CSV.

События — StoredAnomaly стенда (`GET /v1/anomalies`), время — местное пользователя чата.
"""

import csv
import io
from collections import Counter
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .texts import ms_local

SEVERITY = {"CRITICAL": "Критично", "WARNING": "Предупреждение", "INFO": "Инфо"}
RESOLUTION = {"CONFIRMED": "Подтверждено", "FALSE_ALARM": "Ложная тревога"}
EVENT_COLUMNS = ["Начало", "Машина", "Событие", "Важность", "Параметр", "Значение", "Описание", "Решение", "Причина",
                 "Кто решил", "Когда решил", "Источник"]
VEHICLE_COLUMNS = ["Машина", "Событий", "Критичных", "Ждут решения", "Чаще всего"]


def _resolved_at(a: dict, offset: int) -> datetime | None:
    return ms_local(a["resolvedAt"], offset) if a.get("resolvedAt") else None


def event_rows(items: list[dict], offset: int) -> list[list]:
    return [[
        datetime.fromisoformat(a["eventTime"]), a["vehicleName"], a["title"], SEVERITY.get(a["severity"], a["severity"]),
        a.get("parameterCaption") or a.get("parameterName"), a.get("value"), a.get("description"),
        RESOLUTION.get(a.get("resolution") or "", "Ждёт решения"), a.get("falseAlarmReason"), a.get("resolvedBy"),
        _resolved_at(a, offset), a.get("source"),
    ] for a in items]


def vehicle_rows(items: list[dict], vehicles: list[dict]) -> list[list]:
    """По каждой машине пользователя, в том числе без событий; сначала с наибольшим числом событий."""
    per: dict[str, list[dict]] = {v["id"]: [] for v in vehicles}
    names = {v["id"]: v["name"] for v in vehicles}
    for a in items:
        per.setdefault(a["vehicleId"], []).append(a)
        names.setdefault(a["vehicleId"], a["vehicleName"])
    rows = [[names[vid], len(rs), sum(a["severity"] == "CRITICAL" for a in rs),
             sum(not a.get("resolution") for a in rs),
             Counter(a["title"] for a in rs).most_common(1)[0][0] if rs else ""] for vid, rs in per.items()]
    return sorted(rows, key=lambda r: (-r[1], r[0]))


def xlsx(items: list[dict], vehicles: list[dict], offset: int) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "События"
    _sheet(ws, EVENT_COLUMNS, event_rows(items, offset), widths=[17, 28, 30, 15, 22, 10, 50, 16, 20, 16, 17, 18])
    for row in ws.iter_rows(min_row=2):
        for i in (0, 10):
            row[i].number_format = "DD.MM.YYYY HH:MM"
    _sheet(wb.create_sheet("По машинам"), VEHICLE_COLUMNS, vehicle_rows(items, vehicles), widths=[30, 10, 11, 14, 30])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _sheet(ws, header: list[str], rows: list[list], widths: list[int]) -> None:  # noqa: ANN001 — лист openpyxl
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def events_csv(items: list[dict], offset: int) -> bytes:
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(EVENT_COLUMNS)
    for r in event_rows(items, offset):
        w.writerow([_cell(v) for v in r])
    return out.getvalue().encode("utf-8-sig")


def _cell(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return f"{v:%d.%m.%Y %H:%M}"
    if isinstance(v, float):
        return f"{v}".replace(".", ",")
    return str(v)
